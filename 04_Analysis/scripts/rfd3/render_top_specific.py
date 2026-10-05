#!/usr/bin/env python3
"""Highlight the top Brd4ET-over-Brd3ET designs and render both complexes.

Downloads the stored DeltaForge folds, colors hotspots and flexible residues
from each run's YAML, and writes a side-by-side PowerPoint.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

import yaml

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
RUNS = REPO / "design" / "rfd3" / "foundry" / "runs"
RFD3_OUT = REPO / "analysis" / "outputs" / "rfd3"
RANK_CSV = RFD3_OUT / "top_0.5pct_sequences_deltaforge_brd4_over_brd3.csv"
ALIGNMENT = REPO / "staging" / "sequences" / "Brd_ET_aligned.fasta"
OUT_DIR = RFD3_OUT / "top10_brd4et_over_brd3"
PLOT_PATH = RFD3_OUT / "brd4et_vs_brd3et_predicted.png"
PYMOL = Path("/opt/homebrew/bin/pymol")
LIGAND_AI = REPO / "design" / "ligand_ai"
DELTAFORGE_SCRIPTS = HERE.parent / "deltaforge"
VENDOR = REPO / "analysis" / ".vendor"

RESIDUE_RE = re.compile(r"^([A-Za-z]+)(\d+)$")
AVOID_KEYS = ("select_avoid", "select_negative_hotspots")
PNG_WIDTH = 1200
PNG_HEIGHT = 900
TOP_N = 10
MOVIE_THROUGH = 5
MOVIE_FRAMES = 240
MOVIE_FPS = 30


def load_top(path: Path, n: int = TOP_N) -> list[dict[str, str]]:
    """First n rows of the Brd4ET-over-Brd3ET specificity ranking."""
    chosen = []
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            if int(row["specificity_rank"]) <= n:
                chosen.append(row)
    chosen.sort(key=lambda row: int(row["specificity_rank"]))
    if len(chosen) < n:
        raise SystemExit(f"{path} has {len(chosen)} rows with specificity_rank <= {n}")
    return chosen


def load_fasta_alignment(path: Path) -> dict[str, str]:
    records: dict[str, str] = {}
    name: str | None = None
    chunks: list[str] = []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith(">"):
            if name is not None:
                records[name] = "".join(chunks)
            name = line[1:].strip()
            chunks = []
        else:
            chunks.append(line)
    if name is not None:
        records[name] = "".join(chunks)
    return records


def brd4_to_brd3(alignment: dict[str, str]) -> dict[int, int]:
    """Map 1-based Brd4ET residue numbers onto Brd3ET through the alignment."""
    brd4 = alignment["Brd4ET"]
    brd3 = alignment["BRD3ET"]
    if len(brd4) != len(brd3):
        raise ValueError("Brd4ET and BRD3ET alignments differ in length")
    mapping: dict[int, int] = {}
    index4 = 0
    index3 = 0
    for left, right in zip(brd4, brd3):
        if left != "-":
            index4 += 1
        if right != "-":
            index3 += 1
        if left != "-" and right != "-":
            mapping[index4] = index3
    return mapping


def residue_number(key: str) -> int:
    match = RESIDUE_RE.fullmatch(str(key).strip())
    if match is None:
        raise ValueError(f"Unrecognized residue key {key!r}")
    return int(match.group(2))


def is_bkbn(value: Any) -> bool:
    return isinstance(value, str) and value.strip().upper() == "BKBN"


def is_fully_flexible(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, list):
        return len(value) == 0
    text = str(value).strip()
    return text in {"", "[]"}


def load_job_spec(run: str) -> dict[str, Any]:
    path = RUNS / run / "rfd3" / "Brd4ET.yaml"
    if not path.is_file():
        raise FileNotFoundError(path)
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not data:
        raise ValueError(f"No jobs in {path}")
    spec = next(iter(data.values()))
    if not isinstance(spec, dict):
        raise ValueError(f"No job spec in {path}")
    return spec


def atom_text(value: Any) -> str:
    if isinstance(value, list):
        return ",".join(str(item) for item in value)
    return str(value).strip()


def loopy_text(spec: dict[str, Any]) -> str:
    flag = spec.get("is_non_loopy")
    if flag is True:
        return "no"
    if flag is False:
        return "yes"
    return "not set"


def hotspot_text(spec: dict[str, Any]) -> str:
    hotspots = spec.get("select_hotspots") or {}
    if not isinstance(hotspots, dict) or not hotspots:
        return "none"
    parts = []
    for key, value in hotspots.items():
        atoms = atom_text(value)
        parts.append(f"{key} ({atoms})" if atoms else str(key))
    return "; ".join(parts)


def flexible_text(spec: dict[str, Any]) -> str:
    fixed = spec.get("select_fixed_atoms") or {}
    if not isinstance(fixed, dict) or not fixed:
        return "none"
    entries: list[tuple[str, int, str]] = []
    for key, value in fixed.items():
        match = RESIDUE_RE.fullmatch(str(key).strip())
        if match is None:
            raise ValueError(f"Unrecognized residue key {key!r}")
        label = "[]" if is_fully_flexible(value) else "BKBN" if is_bkbn(value) else atom_text(value)
        entries.append((match.group(1), int(match.group(2)), label))
    entries.sort(key=lambda item: (item[0], item[1]))
    ranges: list[tuple[str, int, int, str]] = []
    for chain, number, label in entries:
        if ranges and ranges[-1][0] == chain and ranges[-1][3] == label and number == ranges[-1][2] + 1:
            ranges[-1] = (chain, ranges[-1][1], number, label)
        else:
            ranges.append((chain, number, number, label))
    parts = []
    for chain, start, end, label in ranges:
        span = f"{chain}{start}" if start == end else f"{chain}{start}–{chain}{end}"
        parts.append(f"{span}: {label}")
    return "; ".join(parts)


def design_spec_rows(run: str) -> list[tuple[str, str]]:
    spec = load_job_spec(run)
    return [
        ("Loopy", loopy_text(spec)),
        ("Hotspots", hotspot_text(spec)),
        ("Flexible", flexible_text(spec)),
    ]


def yaml_sets(run: str) -> tuple[list[int], list[int], list[int], list[int]]:
    spec = load_job_spec(run)
    hotspots = [residue_number(key) for key in (spec.get("select_hotspots") or {})]
    fixed = spec.get("select_fixed_atoms") or {}
    flex_full = [residue_number(key) for key, value in fixed.items() if is_fully_flexible(value)]
    flex_bkbn = [residue_number(key) for key, value in fixed.items() if is_bkbn(value)]
    avoid_src: dict[str, Any] = {}
    for key in AVOID_KEYS:
        block = spec.get(key) or {}
        if isinstance(block, dict):
            avoid_src.update(block)
    avoid = [residue_number(key) for key in avoid_src]
    return hotspots, flex_full, flex_bkbn, avoid


def translate(numbers: list[int], receptor: str, mapping: dict[int, int]) -> list[int]:
    if receptor == "Brd4ET":
        return list(numbers)
    translated = []
    for number in numbers:
        hit = mapping.get(number)
        if hit is None:
            print(f"  Brd4 residue {number} has no Brd3ET alignment column", file=sys.stderr)
            continue
        translated.append(hit)
    return translated


def resi_selection(obj: str, numbers: list[int]) -> str:
    if not numbers:
        return "none"
    residues = "+".join(str(number) for number in sorted(set(numbers)))
    return f"{obj} and chain A and resi {residues}"


def color_commands(obj: str, hotspots: list[int], flex_full: list[int], flex_bkbn: list[int], avoid: list[int]) -> list[str]:
    lines = [
        f"cmd.hide('everything', {obj!r})",
        f"cmd.show('cartoon', {obj!r})",
        f"cmd.color('cyan', {obj + ' and chain A'!r})",
        f"cmd.color('purple', {obj + ' and chain B'!r})",
    ]

    def add(name: str, numbers: list[int], extra: list[str]) -> None:
        selection = resi_selection(obj, numbers)
        lines.append(f"cmd.select({name!r}, {selection!r})")
        if numbers:
            lines.extend(extra)

    add(f"{obj}_flex", flex_full, [f"cmd.show('sticks', {obj + '_flex'!r})", f"cmd.color('blue', {obj + '_flex'!r})"])
    add(f"{obj}_avoid", avoid, [f"cmd.show('sticks', {obj + '_avoid'!r})", f"cmd.color('red', {obj + '_avoid'!r})"])
    add(f"{obj}_hot", hotspots, [f"cmd.show('sticks', {obj + '_hot'!r})", f"cmd.color('green', {obj + '_hot'!r})"])
    bkbn = resi_selection(obj, flex_bkbn)
    lines.append(f"cmd.select({obj + '_bkbn'!r}, {bkbn!r})")
    if flex_bkbn:
        lines.append(f"cmd.color('blue', {obj + '_bkbn and name n+ca+c+o'!r})")
    return lines


def write_structures(client: Any, row: dict[str, str], dest: Path) -> tuple[Path, Path]:
    dest.mkdir(parents=True, exist_ok=True)
    written: dict[str, Path] = {}
    for receptor, column in (("Brd4ET", "brd4_job_id"), ("Brd3ET", "brd3_job_id")):
        job_id = (row.get(column) or "").strip()
        if not job_id:
            raise SystemExit(f"{row['model_id']} is missing {column}")
        for fmt in ("pdb", "cif"):
            text = client.get_folding_job_structure(job_id, fmt)
            if not isinstance(text, str) or len(text) < 100:
                raise RuntimeError(f"{job_id} returned no {fmt} structure")
            path = dest / f"{receptor}.{fmt}"
            path.write_text(text if text.endswith("\n") else text + "\n", encoding="utf-8")
            written[receptor if fmt == "pdb" else f"{receptor}_cif"] = path
            print(f"  wrote {path.relative_to(RFD3_OUT)} ({path.stat().st_size} bytes)", flush=True)
    return written["Brd4ET"], written["Brd3ET"]


def ffmpeg_env() -> dict[str, str]:
    """Environment whose PATH can run ffmpeg, including conda copies off PATH."""
    env = os.environ.copy()
    if shutil.which("ffmpeg", path=env.get("PATH", "")):
        return env
    conda = Path.home() / "miniconda3"
    candidates = [conda / "bin" / "ffmpeg"]
    envs = conda / "envs"
    if envs.is_dir():
        candidates.extend(sorted(envs.glob("*/bin/ffmpeg")))
    for binary in candidates:
        if not binary.is_file():
            continue
        probe = subprocess.run(
            [str(binary), "-version"],
            check=False,
            capture_output=True,
            text=True,
        )
        if probe.returncode == 0 and "ffmpeg version" in probe.stdout:
            env["PATH"] = str(binary.parent) + os.pathsep + env.get("PATH", "")
            return env
    raise SystemExit("ffmpeg is required for --movie and was not found")


def render_deck(
    rows: list[dict[str, str]],
    mapping: dict[int, int],
    *,
    movies: bool = False,
) -> list[tuple[dict[str, str], Path, Path]]:
    """Align every complex to the closed-state ET and render one shared camera."""
    snippets = REPO / "analysis" / "pymol" / "snippets"
    panels: list[tuple[dict[str, str], str, Path, Path]] = []
    body = [
        "import sys",
        f"sys.path.insert(0, {str(snippets)!r})",
        "from pymol import cmd",
        "from pymol_render import anchor_et_view, scale_view_to_fit",
        "cmd.reinitialize()",
        "cmd.bg_color('white')",
        "cmd.set('ray_opaque_background', 1)",
        "cmd.set('opaque_background', 1)",
        "cmd.set('cartoon_fancy_helices', 1)",
        "saved = None",
    ]
    for row in rows:
        dest = OUT_DIR / f"{row['specificity_rank']}_{row['model_id']}"
        hotspots, flex_full, flex_bkbn, avoid = yaml_sets(row["run"])
        for label, pdb_name, shifted in (
            ("Brd4ET", "Brd4ET.pdb", False),
            ("Brd3ET", "Brd3ET.pdb", True),
        ):
            pdb = dest / pdb_name
            if not pdb.is_file() or pdb.stat().st_size < 1000:
                raise SystemExit(f"Missing structure {pdb}")
            obj = f"c{row['specificity_rank']}_{label}"
            png = dest / f"{label}.png"
            panels.append((row, obj, png, dest))
            body.append(f"cmd.load({str(pdb.resolve())!r}, {obj!r})")
            numbers = (
                translate(hotspots, "Brd3ET", mapping),
                translate(flex_full, "Brd3ET", mapping),
                translate(flex_bkbn, "Brd3ET", mapping),
                translate(avoid, "Brd3ET", mapping),
            ) if shifted else (hotspots, flex_full, flex_bkbn, avoid)
            body.extend(color_commands(obj, *numbers))
            body.append(f"saved = anchor_et_view({obj!r}, 'A')")
    body.extend(
        [
            "store = {'reference': [], 'complexes': []}",
            "cmd.iterate_state(1, 'receptor', 'reference.append((x,y,z))', space=store)",
            "cmd.iterate_state(1, 'not receptor', 'complexes.append((x,y,z))', space=store)",
            "fitted = scale_view_to_fit(",
            "    saved,",
            "    store['complexes'],",
            f"    width={PNG_WIDTH},",
            f"    height={PNG_HEIGHT},",
            "    reference_points=store['reference'],",
            ")",
            "print('ET camera pull-back %.4f' % (fitted[11] / saved[11]))",
            "cmd.set_view(fitted)",
        ]
    )
    if movies:
        body.extend(
            [
                "import shutil",
                "import subprocess",
                "from pathlib import Path",
            ]
        )
    for _row, obj, png, dest in panels:
        body.extend(
            [
                "cmd.disable('all')",
                f"cmd.enable({obj!r})",
                "cmd.disable('receptor')",
                "cmd.set_view(fitted)",
                f"cmd.png({str(png.resolve())!r}, width={PNG_WIDTH}, height={PNG_HEIGHT}, ray=0, dpi=120)",
            ]
        )
        if movies and int(_row["specificity_rank"]) <= MOVIE_THROUGH:
            mp4 = dest / png.name.replace(".png", ".mp4")
            frames = dest / f"_frames_{png.stem}"
            # Headless PyMOL leaves movie.produce/mpng frames blank, so draw each
            # turn with cmd.png and stitch. Frame 1 stays on the anchored camera.
            body.extend(
                [
                    "cmd.set_view(fitted)",
                    f"print({obj!r}, flush=True)",
                    f"open({str((OUT_DIR / '_movie_progress.txt').resolve())!r}, 'a').write({obj!r} + '\\n')",
                    f"frames = Path({str(frames.resolve())!r})",
                    "if frames.exists():",
                    "    shutil.rmtree(frames)",
                    "frames.mkdir()",
                    f"for frame in range(1, {MOVIE_FRAMES + 1}):",
                    "    if frame > 1:",
                    "        cmd.turn('y', -1.506)",
                    "    cmd.png(",
                    "        str(frames / ('frame%04d.png' % frame)),",
                    f"        width={PNG_WIDTH},",
                    f"        height={PNG_HEIGHT},",
                    "        ray=0,",
                    "        dpi=120,",
                    "    )",
                    "subprocess.run(",
                    "    [",
                    "        'ffmpeg', '-y', '-loglevel', 'error',",
                    f"        '-framerate', '{MOVIE_FPS}',",
                    "        '-start_number', '1',",
                    "        '-i', str(frames / 'frame%04d.png'),",
                    "        '-c:v', 'libx264',",
                    "        '-pix_fmt', 'yuv420p',",
                    f"        {str(mp4.resolve())!r},",
                    "    ],",
                    "    check=True,",
                    ")",
                    "shutil.rmtree(frames)",
                    "cmd.set_view(fitted)",
                ]
            )
    body.append("cmd.refresh()")
    progress = OUT_DIR / "_movie_progress.txt"
    if progress.exists():
        progress.unlink()
    pml = OUT_DIR / "_render.pml"
    pml.write_text("python\n" + "\n".join(body) + "\npython end\nquit\n", encoding="utf-8")
    completed = subprocess.run(
        [str(PYMOL), "-c", "-q", str(pml)],
        check=False,
        capture_output=True,
        text=True,
        env=ffmpeg_env() if movies else None,
    )
    log = completed.stdout + completed.stderr
    missing = [png for _row, _obj, png, _dest in panels if not png.is_file()]
    if movies:
        missing.extend(
            dest / png.name.replace(".png", ".mp4")
            for row, _obj, png, dest in panels
            if int(row["specificity_rank"]) <= MOVIE_THROUGH
            and not (dest / png.name.replace(".png", ".mp4")).is_file()
        )
    if completed.returncode != 0 or "Error:" in log or missing:
        sys.stderr.write(log[-4000:])
        raise RuntimeError(f"PyMOL failed while rendering the shared ET camera (exit {completed.returncode})")
    pullback = [line for line in log.splitlines() if "ET camera pull-back" in line]
    if pullback:
        print(pullback[-1], flush=True)
    if progress.exists():
        progress.unlink()
    slides: list[tuple[dict[str, str], Path, Path]] = []
    for row in rows:
        dest = OUT_DIR / f"{row['specificity_rank']}_{row['model_id']}"
        slides.append((row, dest / "Brd4ET.png", dest / "Brd3ET.png"))
    return slides


def _write_lines(text_frame, lines: list[tuple[str, str, str]], size: int) -> None:
    """Write label/value lines. The third item is the value font name."""
    from pptx.util import Pt

    text_frame.word_wrap = True
    for index, (label, value, font_name) in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        label_run = paragraph.add_run()
        label_run.text = f"{label}  "
        label_run.font.size = Pt(size)
        label_run.font.bold = True
        label_run.font.name = "Calibri"
        value_run = paragraph.add_run()
        value_run.text = value
        value_run.font.size = Pt(size)
        value_run.font.bold = False
        value_run.font.name = font_name


def _metric(row: dict[str, str], prefix: str, field: str, digits: int) -> str:
    text = (row.get(f"{prefix}_{field}") or "").strip()
    try:
        return f"{float(text):.{digits}f}"
    except ValueError:
        return text


def png_size(path: Path) -> tuple[int, int]:
    import struct

    with path.open("rb") as handle:
        header = handle.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"{path} is not a PNG")
    width, height = struct.unpack(">II", header[16:24])
    return width, height


def add_plot_slide(prs, n_designs: int) -> None:
    """First slide: the six-panel scatter whose red labels are the later slide numbers."""
    from pptx.util import Emu, Inches, Pt

    if not PLOT_PATH.is_file():
        raise FileNotFoundError(PLOT_PATH)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.35), Inches(0.08), Inches(12.6), Inches(0.36))
    title_run = title.text_frame.paragraphs[0].add_run()
    title_run.text = (
        f"Top {n_designs} Brd4ET-specific  ·  red labels 1–{n_designs} are the following slides"
    )
    title_run.font.size = Pt(16)
    title_run.font.bold = True

    pixel_w, pixel_h = png_size(PLOT_PATH)
    max_width = Inches(12.6)
    max_height = Inches(6.85)
    width = max_width
    height = Emu(int(width * pixel_h / pixel_w))
    if height > max_height:
        height = max_height
        width = Emu(int(height * pixel_w / pixel_h))
    left = Emu(int((prs.slide_width - width) / 2))
    slide.shapes.add_picture(str(PLOT_PATH), left, Inches(0.48), width, height)


def add_slide(prs, row: dict[str, str], brd4_png: Path, brd3_png: Path) -> None:
    from pptx.util import Inches, Pt

    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.35), Inches(0.08), Inches(12.6), Inches(0.32))
    title_run = title.text_frame.paragraphs[0].add_run()
    title_run.text = f"{row['specificity_rank']}.  {row['model_id']}   ({row['run']})"
    title_run.font.size = Pt(18)
    title_run.font.bold = True

    for left, label in ((0.35, "Brd4ET"), (6.85, "Brd3ET")):
        box = slide.shapes.add_textbox(Inches(left), Inches(0.40), Inches(6.1), Inches(0.24))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = label
        run.font.size = Pt(14)
        run.font.bold = True

    picture_top = 0.64
    picture_height = 3.75
    for left, png in ((0.35, brd4_png), (6.85, brd3_png)):
        slide.shapes.add_picture(
            str(png), Inches(left), Inches(picture_top), Inches(6.1), Inches(picture_height)
        )

    binder = (row.get("binder_sequence") or "").strip()
    if not binder:
        raise SystemExit(f"{row['model_id']} is missing binder_sequence")
    score_lines = [
        ("Kd (nM)", "kd_nm", 2),
        ("ΔG", "delta_g", 2),
        ("ipTM", "iptm", 3),
        ("pTM", "ptm", 3),
        ("Peptide ipSAE", "peptide_ipsae", 3),
        ("Mean pLDDT", "mean_plddt", 1),
    ]
    for left, prefix in ((0.35, "brd4"), (6.85, "brd3")):
        score_box = slide.shapes.add_textbox(Inches(left), Inches(4.46), Inches(6.1), Inches(1.32))
        _write_lines(
            score_box.text_frame,
            [(label, _metric(row, prefix, field, digits), "Calibri") for label, field, digits in score_lines],
            13,
        )

    spec_rows = design_spec_rows(row["run"])
    shared = slide.shapes.add_textbox(Inches(0.35), Inches(5.82), Inches(12.6), Inches(1.56))
    _write_lines(
        shared.text_frame,
        [
            ("Designed chain", binder, "Consolas"),
            *[(label, text, "Calibri") for label, text in spec_rows],
            ("Specificity rank", row["specificity_rank"], "Calibri"),
            ("log10 Kd ratio", f"{float(row['log10_kd_ratio']):.3f}", "Calibri"),
        ],
        12,
    )


def build_deck(slides: list[tuple[dict[str, str], Path, Path]]) -> Path:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        if str(VENDOR) not in sys.path and VENDOR.is_dir():
            sys.path.insert(0, str(VENDOR))
        from pptx import Presentation
        from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_plot_slide(prs, len(slides))
    for row, brd4_png, brd3_png in slides:
        add_slide(prs, row, brd4_png, brd3_png)
    out = OUT_DIR / "top10_brd4et_over_brd3.pptx"
    prs.save(out)
    print(f"Wrote {out}")
    return out


def highlight_plot() -> None:
    if str(DELTAFORGE_SCRIPTS) not in sys.path:
        sys.path.insert(0, str(DELTAFORGE_SCRIPTS))
    from plot_brd4_vs_brd3_deltaforge import (  # noqa: WPS433
        INPUT_PATH,
        OUTPUT_PATH,
        _save_figure,
        load_unique_rows,
    )

    rows = load_unique_rows(INPUT_PATH)
    slides = load_top(RANK_CSV)
    top = [(row["specificity_rank"], row["sequence"]) for row in slides]
    known = {row["sequence"] for row in rows}
    missing = [sequence for _rank, sequence in top if sequence not in known]
    if missing:
        raise SystemExit(f"{len(missing)} highlighted sequences are missing from {INPUT_PATH}")
    _save_figure(
        rows,
        OUTPUT_PATH,
        top,
        highlight_label=f"Top {TOP_N} Brd4ET-specific",
    )


def structures_ready(dest: Path) -> bool:
    return all((dest / name).is_file() and (dest / name).stat().st_size > 1000 for name in ("Brd4ET.pdb", "Brd3ET.pdb"))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--movie",
        action="store_true",
        help="Write an 8-second Y-roll mp4 beside the PNG for ranks 1-5; the deck stays stills",
    )
    parser.add_argument(
        "--slides-only",
        action="store_true",
        help="Rewrite the PowerPoint from existing PNGs",
    )
    args = parser.parse_args()
    if args.slides_only:
        highlight_plot()
        rows = load_top(RANK_CSV)
        slides: list[tuple[dict[str, str], Path, Path]] = []
        for row in rows:
            dest = OUT_DIR / f"{row['specificity_rank']}_{row['model_id']}"
            pngs = (dest / "Brd4ET.png", dest / "Brd3ET.png")
            missing = [path for path in pngs if not path.is_file()]
            if missing:
                print("Missing media:\n" + "\n".join(str(path) for path in missing), file=sys.stderr)
                return 2
            slides.append((row, pngs[0], pngs[1]))
        build_deck(slides)
        return 0
    if not PYMOL.is_file():
        print(f"PyMOL not found: {PYMOL}", file=sys.stderr)
        return 2
    if args.movie:
        ffmpeg_env()
    highlight_plot()
    rows = load_top(RANK_CSV)
    mapping = brd4_to_brd3(load_fasta_alignment(ALIGNMENT))
    client = None
    for row in rows:
        dest = OUT_DIR / f"{row['specificity_rank']}_{row['model_id']}"
        if structures_ready(dest):
            continue
        if client is None:
            sys.path.insert(0, str(LIGAND_AI))
            if not (LIGAND_AI / "src" / "ligandai_local").is_dir():
                print(f"LigandAI package not found at {LIGAND_AI}", file=sys.stderr)
                return 2
            from src.ligandai_local.client import LigandAIClient  # noqa: WPS433

            client = LigandAIClient()
        print(f"{row['specificity_rank']} {row['model_id']}", flush=True)
        write_structures(client, row, dest)
    build_deck(render_deck(rows, mapping, movies=args.movie))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
