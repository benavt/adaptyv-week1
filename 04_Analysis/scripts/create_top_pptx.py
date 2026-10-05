#!/usr/bin/env python3
"""One PowerPoint per top-0.5% PyMOL session, with one slide per binder.

Each slide shows that model alone (other objects disabled), plus ipTM, pTM,
TM, RMSD, and the designed sequence.
"""

from __future__ import annotations

import argparse
import csv
import subprocess
import sys
import textwrap
from collections import defaultdict
from pathlib import Path

from create_top_pse import build_model

ANALYSIS = Path(__file__).resolve().parents[1]
FOUNDRY = ANALYSIS.parent / "design" / "rfd3" / "foundry"
RUNS = FOUNDRY / "runs"
DEFAULT_CSV = FOUNDRY / "reports" / "top_0.5pct_sequences.csv"
PYMOL = Path("/opt/homebrew/bin/pymol")
SNIPPETS = ANALYSIS / "pymol" / "snippets"
sys.path.insert(0, str(SNIPPETS))
from pymol_render import ET_REFERENCE, ET_VIEW  # noqa: E402

PNG_WIDTH = 1200
PNG_HEIGHT = 900


def load_models(csv_path: Path, only: str | None) -> dict[str, list[dict[str, str]]]:
    """Unique models per run, in the same first-seen order as the PyMOL session."""
    by_run: dict[str, list[dict[str, str]]] = defaultdict(list)
    seen: dict[str, set[str]] = defaultdict(set)
    with csv_path.open(newline="") as handle:
        for row in csv.DictReader(handle):
            run = row["run"]
            if only and run != only:
                continue
            model_id = row["model_id"]
            if model_id in seen[run]:
                continue
            seen[run].add(model_id)
            by_run[run].append(row)
    return dict(by_run)


def load_ptm(run: str) -> dict[str, str]:
    scores = RUNS / run / "scores.csv"
    ptm: dict[str, str] = {}
    with scores.open(newline="") as handle:
        for row in csv.DictReader(handle):
            model_id = Path(row["rf3_model"]).name.removesuffix("_model.cif")
            ptm[model_id] = f"{float(row['ptm']):.4f}"
    return ptm


def target_chains(run: str, model_ids: list[str]) -> dict[str, str]:
    """Output chain that carries the job-rendering receptor (input chain A)."""
    run_dir = RUNS / run
    return {model_id: build_model(run_dir, model_id).target_chain for model_id in model_ids}


def render_pngs(run: str, model_ids: list[str], png_dir: Path) -> None:
    pse = RUNS / run / "run_display" / f"{run}_top_0.5pct.pse"
    if not pse.is_file():
        raise FileNotFoundError(pse)
    use_job_view = (RUNS / run / "run_display").glob("*_job_rendering.png")
    use_job_view = next(use_job_view, None) is not None
    chains = target_chains(run, model_ids) if use_job_view else {}
    if use_job_view and (not ET_REFERENCE.is_file() or not ET_VIEW.is_file()):
        raise FileNotFoundError("Closed-state ET structure or camera view is missing")
    png_dir.mkdir(parents=True, exist_ok=True)
    lines = [
        "import sys",
        f"sys.path.insert(0, {str(SNIPPETS.resolve())!r})",
        "from pymol import cmd",
        "from pymol_render import anchor_et_view",
        "cmd.reinitialize()",
        f"cmd.load({str(pse.resolve())!r})",
    ]
    lines.extend(
        [
            "cmd.bg_color('white')",
            "cmd.set('ray_opaque_background', 1)",
            "cmd.set('opaque_background', 1)",
            "names = cmd.get_object_list()",
            f"wanted = {model_ids!r}",
            f"chains = {chains!r}",
            "missing = [name for name in wanted if name not in names]",
            "if missing:",
            "    raise SystemExit('PSE is missing objects: ' + ', '.join(missing))",
            "for name in wanted:",
        ]
    )
    if use_job_view:
        lines.extend(
            [
                "    anchor_et_view(name, chains[name])",
                "    cmd.disable('all')",
                "    cmd.enable(name)",
                "    cmd.disable('receptor')",
            ]
        )
    else:
        lines.extend(
            [
                "    cmd.disable('all')",
                "    cmd.enable(name)",
                "    cmd.zoom(name, 2, complete=1)",
            ]
        )
    lines.extend(
        [
            f"    path = {str(png_dir.resolve())!r} + '/' + name + '.png'",
            f"    cmd.png(path, width={PNG_WIDTH}, height={PNG_HEIGHT}, ray=0, dpi=100)",
            "    cmd.refresh()",
        ]
    )
    pml = png_dir / "_render.pml"
    pml.write_text("python\n" + "\n".join(lines) + "\npython end\nquit\n", encoding="utf-8")
    completed = subprocess.run(
        [str(PYMOL), "-c", "-q", str(pml)],
        check=False,
        capture_output=True,
        text=True,
    )
    log = completed.stdout + completed.stderr
    if completed.returncode != 0 or "Selector-Error" in log or "Error:" in log:
        sys.stderr.write(log[-4000:])
        raise RuntimeError(
            f"PyMOL failed while rendering {run} (exit {completed.returncode})"
        )
    missing = [model_id for model_id in model_ids if not (png_dir / f"{model_id}.png").is_file()]
    if missing:
        sys.stderr.write(completed.stdout[-4000:])
        sys.stderr.write(completed.stderr[-4000:])
        raise RuntimeError(f"PyMOL did not write {len(missing)} PNGs for {run}")


def add_slide(prs, row: dict[str, str], ptm: str, png_path: Path) -> None:
    from pptx.util import Inches, Pt

    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.45))
    title_frame = title.text_frame
    title_frame.word_wrap = False
    title_run = title_frame.paragraphs[0].add_run()
    title_run.text = row["model_id"]
    title_run.font.size = Pt(18)
    title_run.font.bold = True

    slide.shapes.add_picture(str(png_path), Inches(0.3), Inches(0.8), Inches(7.4), Inches(5.55))

    metrics = slide.shapes.add_textbox(Inches(7.9), Inches(0.9), Inches(5.0), Inches(1.6))
    metrics_frame = metrics.text_frame
    metrics_frame.word_wrap = True
    lines = [
        f"ipTM  {row['iptm']}",
        f"pTM   {ptm}",
        f"TM    {row['tm']}",
        f"RMSD  {row['rmsd']}",
    ]
    metrics_frame.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        paragraph = metrics_frame.add_paragraph()
        paragraph.text = line
    for paragraph in metrics_frame.paragraphs:
        paragraph.font.size = Pt(20)
        paragraph.font.name = "Menlo"

    seq_box = slide.shapes.add_textbox(Inches(7.9), Inches(2.8), Inches(5.0), Inches(4.2))
    seq_frame = seq_box.text_frame
    seq_frame.word_wrap = True
    label = seq_frame.paragraphs[0]
    label.text = "Designed sequence"
    label.font.size = Pt(14)
    label.font.bold = True
    body = seq_frame.add_paragraph()
    body.text = "\n".join(textwrap.wrap(row["sequence"], width=42))
    body.font.size = Pt(11)
    body.font.name = "Menlo"


def build_deck(run: str, rows: list[dict[str, str]]) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    ptm_by_id = load_ptm(run)
    png_dir = RUNS / run / "run_display" / "top_0.5pct_png"
    model_ids = [row["model_id"] for row in rows]
    render_pngs(run, model_ids, png_dir)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for row in rows:
        model_id = row["model_id"]
        if model_id not in ptm_by_id:
            raise KeyError(f"{run} scores.csv has no pTM for {model_id}")
        add_slide(prs, row, ptm_by_id[model_id], png_dir / f"{model_id}.png")
    out = RUNS / run / "run_display" / f"{run}_top_0.5pct.pptx"
    prs.save(out)
    print(f"{run}: {len(rows)} slides -> {out}")
    return out


def ensure_pptx() -> None:
    vendor = ANALYSIS / ".vendor"
    if vendor.is_dir():
        sys.path.insert(0, str(vendor))
    try:
        import pptx  # noqa: F401
    except ImportError:
        vendor.mkdir(exist_ok=True)
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--target", str(vendor), "python-pptx"]
        )
        if str(vendor) not in sys.path:
            sys.path.insert(0, str(vendor))
        import pptx  # noqa: F401


def main() -> None:
    ensure_pptx()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    parser.add_argument("--run", help="Build only this run directory name")
    args = parser.parse_args()
    runs = load_models(args.csv, args.run)
    if not runs:
        raise SystemExit(f"No models found in {args.csv}")
    for run, rows in runs.items():
        build_deck(run, rows)


if __name__ == "__main__":
    main()
