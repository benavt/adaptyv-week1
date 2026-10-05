#!/usr/bin/env python3
import csv,json,subprocess,sys
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
from molecular_colors import color_commands, COLOR_KEY
from Bio.PDB import MMCIFParser
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
ROOT=Path(__file__).resolve().parents[2]; IN=ROOT/'03_Filtering/Refolding/Putative_DD1_EGFR_D3'; OUT=ROOT/'03_Filtering/PyMol_top20/intersection_both_species'; REND=OUT/'renders'; PML=OUT/'pml'; OUT.mkdir(parents=True,exist_ok=True); REND.mkdir(exist_ok=True); PML.mkdir(exist_ok=True)
DESIGNS=['DD1_062','DD1_073','DD1_115','DD1_245']; VIEW=ROOT/'03_Filtering/PyMol_top20/view_capture'; DF=ROOT/'03_Filtering/DeltaForge/out/intersection_both_species/results.csv'; HIS=ROOT/'03_Filtering/JustHIpKA/intersection_both_species/hispka_results.csv'; LONG=ROOT/'03_Filtering/Boltz2_metrics/out/boltz2_dd1_metrics_long.csv'
csv.field_size_limit(sys.maxsize)
def cif_for(d,s): return next(IN.glob(f'{d}_{s}_*/{d}_{s}_*_model_0.cif'))
def txt(slide,text,x,y,w,h,size=14,bold=False,align=PP_ALIGN.LEFT,color=(30,40,50)):
 if isinstance(align, tuple): color,align=align,PP_ALIGN.LEFT
 b=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*color); return b
def make_pml(d,s):
 cif=cif_for(d,s); view=json.loads((VIEW/('human_view.json' if s=='Human' else 'mouse_view.json')).read_text()); ref=ROOT/'01_Staging/EGFR'/('Human_EGFR_seg_medoid.pdb' if s=='Human' else 'Mouse_EGFR_seg_medoid.pdb'); out=REND/f'{d}_{s}.png'; p=PML/f'{d}_{s}.pml'
 p.write_text(f'''load {cif}, complex\nload {ref}, reference\nhide everything, reference\nshow cartoon, complex\nshow sticks, complex and chain B and (resi 37 or resi 100) and resn HIS\nshow sticks, complex and chain A and resn ASP\n{color_commands(s)}\nalign complex and chain B and name CA, reference and chain A and name CA\nset_view ({', '.join(str(v) for v in view)})\nbg_color white\nset orthoscopic, on\nset ray_opaque_background, off\nviewport 1400,900\nray 1400,900\npng {out}, dpi=180\nquit\n'''); subprocess.run(['/opt/homebrew/bin/pymol','-cq',str(p)],check=True); return out
def f(v):
 try:return f'{float(v):.3f}'
 except:return '—'
def main():
 dfr={(r['design_id'],r['species']):r for r in csv.DictReader(DF.open())}; his={(r['design_id'],r['species'],r['residue']):r for r in csv.DictReader(HIS.open())}; metrics={}
 for r in csv.DictReader(LONG.open()):
  for d in DESIGNS:
   if f'DD1_{int(d.split("_")[1]):03d}_' in r['boltz_dir'] and r['target'].lower() in ('human','mouse'): metrics[(d,r['target'].title())]=r
 prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
 s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(248,250,252); txt(s,'DD1 ASP–HIS proximity intersection',.45,.25,12.4,.4,24,True); txt(s,'Subset definition: at least one DD1 ASP within 5 Å of both EGFR HIS37 and HIS100 in the Human and Mouse refolded complexes. Four designs passed this intersection.',.55,.9,12.2,.7,15,False); txt(s,'Designs: DD1_062, DD1_073, DD1_115, DD1_245',.55,2.0,12,.35,20,True,(50,70,90)); txt(s,'Scores shown below are remote DeltaForge predictions and JustHISpKa predictions for the EGFR target chains.',.55,6.75,12,.25,11,False,(90,100,110))
 for d in DESIGNS:
  s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(248,250,252); txt(s,f'{d} | EGFR D3 interface',.45,.18,12.4,.36,22,True); txt(s,'Human EGFR D3',.55,.62,5.9,.25,15,True,(35,55,75),PP_ALIGN.CENTER); txt(s,'Mouse EGFR D3',6.88,.62,5.9,.25,15,True,(35,55,75),PP_ALIGN.CENTER)
  for sp,x in [('Human',.35),('Mouse',6.68)]: s.shapes.add_picture(str(make_pml(d,sp)),Inches(x),Inches(.92),width=Inches(6.3),height=Inches(4.12))
  table=s.shapes.add_table(7,3,Inches(.55),Inches(5.45),Inches(12.2),Inches(1.55)).table; table.columns[0].width=Inches(2.8); table.columns[1].width=Inches(4.7); table.columns[2].width=Inches(4.7)
  vals=[('Metric','Human EGFR D3','Mouse EGFR D3'),('ipSAE',f(metrics[(d,'Human')]['ipSAE']),f(metrics[(d,'Mouse')]['ipSAE'])),('TM score',f(metrics[(d,'Human')]['tm_score']),f(metrics[(d,'Mouse')]['tm_score'])),('DeltaForge ΔG (kcal/mol)',f(dfr[(d,'Human')].get('delta_g')),f(dfr[(d,'Mouse')].get('delta_g'))),('DeltaForge Kd (nM)',f(dfr[(d,'Human')].get('kd_nm')),f(dfr[(d,'Mouse')].get('kd_nm'))),('HIS37 / HIS100 pKa',f(his[(d,'Human','37')]['pKa'])+' / '+f(his[(d,'Human','100')]['pKa']),f(his[(d,'Mouse','37')]['pKa'])+' / '+f(his[(d,'Mouse','100')]['pKa']))]
  for i,row in enumerate(vals):
   for j,val in enumerate(row):
    c=table.cell(i,j); c.text=val; c.vertical_anchor=MSO_ANCHOR.MIDDLE; c.fill.solid(); c.fill.fore_color.rgb=RGBColor(225,232,239) if i==0 else RGBColor(248,250,252)
    for p in c.text_frame.paragraphs:
     p.alignment=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER
     for r in p.runs:r.font.name='Arial'; r.font.size=Pt(8); r.font.bold=(i==0 or j==0); r.font.color.rgb=RGBColor(30,40,50)
  txt(s,COLOR_KEY,.45,5.25,12.4,.16,7.5,False,(90,100,110),PP_ALIGN.CENTER)
 prs.save(OUT/'DD1_ASP_both_species_intersection.pptx'); print('saved',OUT/'DD1_ASP_both_species_intersection.pptx')
if __name__=='__main__':main()
