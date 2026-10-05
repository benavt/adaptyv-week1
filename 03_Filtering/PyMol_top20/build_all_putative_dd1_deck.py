#!/usr/bin/env python3
"""Build a complete paired Human/Mouse metrics deck for Putative_DD1."""
import csv, re, shutil, subprocess
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
from molecular_colors import COLOR_KEY
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[2]
METRICS = ROOT / '03_Filtering/Boltz2_metrics/out/boltz2_dd1_metrics_long.csv'
SEQ = ROOT / '03_Filtering/RFD3_filtering/out/dd1_asp_near_both_his_sequences.csv'
PKA = ROOT / '03_Filtering/JustHIpKA/putative_DD1_EGFR_D3'
DOCKQ_TM = ROOT / '03_Filtering/PyMol_top20/out/dockq_tm_all_designs.csv'
OUT = ROOT / '03_Filtering/PyMol_top20/Putative_DD1_EGFR_metrics.pptx'
IMG = ROOT / '03_Filtering/PyMol_top20/renders'
SCATTER_SVG = ROOT / '03_Filtering/Boltz2_metrics/out/scatter_ipSAE.svg'
SCATTER_PNG = ROOT / '03_Filtering/PyMol_top20/putative_dd1_ipSAE_scatter.png'

def safe(v): return re.sub(r'[^A-Za-z0-9_.-]+', '_', v)
def short_model_name(v):
    return re.sub(r'^(Human|Mouse)_EGFR_(Human|Mouse)_EGFR_', r'\1_', v)
def val(v, digits=3):
    try: return f'{float(v):.{digits}f}'
    except (TypeError, ValueError): return 'N/A'
def pka(design, target, sid):
    path = PKA / f'{safe(design)}_{target}' / 'predictions_A_HIS37_HIS100' / f'HIS{sid}.txt'
    if not path.exists(): return 'N/A'
    m = re.search(r'PKA=([0-9.]+)', path.read_text())
    return m.group(1) if m else 'N/A'
def text(slide, content, x, y, w, h, size=12, bold=False, color=(30,40,50), align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(content); r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*color)
    return box
def table(slide, rows, x, y, w, h, widths=None, font_size=8):
    t=slide.shapes.add_table(len(rows),len(rows[0]),Inches(x),Inches(y),Inches(w),Inches(h)).table
    widths = widths or [w / len(rows[0])] * len(rows[0])
    for col, width in zip(t.columns, widths): col.width=Inches(width)
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=str(v); c.vertical_anchor=MSO_ANCHOR.MIDDLE; c.fill.solid(); c.fill.fore_color.rgb=RGBColor(225,232,239) if i==0 else RGBColor(248,250,252)
            for p in c.text_frame.paragraphs:
                p.alignment=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER
                for r in p.runs: r.font.name='Arial'; r.font.size=Pt(font_size); r.font.bold=(i==0 or j==0); r.font.color.rgb=RGBColor(30,40,50)
def main():
    metrics=list(csv.DictReader(METRICS.open())); seq={r['design_id']:r for r in csv.DictReader(SEQ.open())}
    dockq_tm={r['design_id']:r for r in csv.DictReader(DOCKQ_TM.open())}
    by={(r['design_id'],r['target']):r for r in metrics}
    def mean_ipsae(d):
        try: return (float(by[(d,'Human')]['ipSAE']) + float(by[(d,'Mouse')]['ipSAE'])) / 2
        except (KeyError, TypeError, ValueError): return float('-inf')
    designs=sorted(seq, key=lambda d: (-mean_ipsae(d), d))
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    # Opening slide: the project-standard paired ipSAE scatterplot.
    if SCATTER_SVG.exists():
        subprocess.run(['rsvg-convert','-w','1400','-o',str(SCATTER_PNG),str(SCATTER_SVG)], check=True)
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(248,250,252)
    text(s,'Putative DD1 designs: Human vs Mouse EGFR ipSAE',.45,.2,12.4,.35,22,True)
    text(s,'Each point is one design. Designs in the deck are ordered by mean ipSAE across the two EGFR species, highest first.',.45,.62,12.4,.25,11,False,(70,80,90))
    if SCATTER_PNG.exists(): s.shapes.add_picture(str(SCATTER_PNG), Inches(3.0), Inches(1.05), width=Inches(7.4))
    text(s,'The dashed diagonal marks equal Human and Mouse ipSAE. See the source Boltz2 paired metrics for exact values.',.45,7.12,12.4,.2,8,False,(90,100,110),PP_ALIGN.CENTER)
    # Overview slides keep every design visible in the deck index.
    for start in range(0,len(designs),30):
        s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(248,250,252)
        page=start//30+1; text(s, f'Putative DD1 designs and paired EGFR metrics | index {page}', .45,.2,12.4,.35,21,True)
        rows=[('DD1 sequence','Model','Human ipSAE','Mouse ipSAE')]
        for d in designs[start:start+30]: rows.append((seq[d]['binder_sequence'],short_model_name(d),val(by[(d,'Human')].get('ipSAE')),val(by[(d,'Mouse')].get('ipSAE'))))
        table(s,rows,.35,.75,12.63,6.55,widths=[7.70,2.00,1.45,1.48],font_size=6.2)
    for n,d in enumerate(designs,1):
        s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(248,250,252)
        text(s,f'{n:03d}  {d} | Human and Mouse EGFR D3',.45,.18,12.4,.35,21,True)
        text(s,'Human EGFR D3',.55,.62,5.9,.25,15,True,(35,55,75),PP_ALIGN.CENTER); text(s,'Mouse EGFR D3',6.88,.62,5.9,.25,15,True,(35,55,75),PP_ALIGN.CENTER)
        # Fixed-view PyMOL renders, sized to preserve the 900x650 source ratio.
        human_png = IMG / f'{d}_Human.png'
        mouse_png = IMG / f'{d}_Mouse.png'
        if human_png.exists(): s.shapes.add_picture(str(human_png), Inches(1.65), Inches(.92), height=Inches(2.72))
        else: text(s,'Human render unavailable',.55,1.05,5.9,2.5,12,False,(110,120,130),PP_ALIGN.CENTER)
        if mouse_png.exists(): s.shapes.add_picture(str(mouse_png), Inches(8.00), Inches(.92), height=Inches(2.72))
        else: text(s,'Mouse render unavailable',6.88,1.05,5.9,2.5,12,False,(110,120,130),PP_ALIGN.CENTER)
        text(s,COLOR_KEY,.45,3.75,12.4,.2,7.5,False,(90,100,110),PP_ALIGN.CENTER)
        h=by[(d,'Human')]; m=by[(d,'Mouse')]
        q=dockq_tm[d]
        text(s, f"DockQ Human→Mouse: {float(q['dockq']):.3f} | iRMSD: {float(q['dockq_irmsd']):.2f} Å | LRMSD: {float(q['dockq_lrmsd']):.2f} Å | Fnat: {float(q['dockq_fnat']):.3f} | TM Mouse→Human: {float(q['tm_score_mouse_vs_human']):.3f} | RMSD: {float(q['tm_rmsd_mouse_vs_human']):.2f} Å", .55, 3.93, 12.2, .2, 8.2, True, (192, 0, 0), PP_ALIGN.CENTER)
        rows=[('Metric','Human EGFR D3','Mouse EGFR D3'),
          ('ipSAE',val(h.get('ipSAE')),val(m.get('ipSAE'))),('ipTM',val(h.get('ipTM_af')),val(m.get('ipTM_af'))),
          ('Complex pLDDT',val(h.get('complex_plddt')),val(m.get('complex_plddt'))),('pDockQ',val(h.get('pDockQ')),val(m.get('pDockQ'))),
          ('Interface mean PAE',val(h.get('ipdae_mean_pae')),val(m.get('ipdae_mean_pae'))),('TM score to RFD3',val(h.get('tm_score')),val(m.get('tm_score'))),
          ('TM RMSD (Å)',val(h.get('tm_rmsd')),val(m.get('tm_rmsd'))),('HIS37 pKa',pka(d,'Human',37),pka(d,'Mouse',37)),('HIS100 pKa',pka(d,'Human',100),pka(d,'Mouse',100))]
        table(s,rows,.55,4.18,12.2,2.47)
        text(s,'Binder sequence: '+seq[d]['binder_sequence'],.55,6.8,12.1,.18,7,False,(70,80,90))
        text(s,'Metrics sourced from Boltz2 paired analysis. Red line: DockQ Human→Mouse with fixed AB:AB chain mapping and whole-complex TM Mouse→Human normalized by the Human model; pKa values sourced from JustHISpKa.',.45,7.12,12.4,.2,7.5,False,(90,100,110),PP_ALIGN.CENTER)
    OUT.parent.mkdir(parents=True,exist_ok=True); prs.save(OUT); print(OUT)
if __name__=='__main__': main()
